#!/usr/bin/env python3
"""
AWS backup cost report.

Produces an Excel workbook showing what this AWS account actually pays for
backup today, plus an inventory that explains the bill and surfaces storage
that has been forgotten about.

Design rules, deliberately:

* **Read only.** Every API call is a List/Describe/Get. Nothing is created,
  modified or deleted.
* **Nothing leaves your environment.** The script talks to AWS and writes a
  local .xlsx. There is no telemetry and no upload.
* **Real billed dollars only.** Every figure in the money columns comes from
  Cost Explorer, i.e. what AWS actually charged. The inventory sheets report
  counts, sizes and ages, and are never priced — estimating them would mix
  modelled numbers into a document whose whole value is that it does not.

Usage:
    python aws_backup_cost_report.py [--profile NAME] [--months 12] [--output FILE]
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import sys
from collections import defaultdict
from typing import Any, Dict, List, Optional, Sequence, Tuple

try:
    import boto3
    from botocore.config import Config
    from botocore.exceptions import BotoCoreError, ClientError, NoCredentialsError
except ImportError:  # pragma: no cover - import guard for a friendlier message
    sys.exit("boto3 is required.  pip install -r requirements.txt")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required.  pip install -r requirements.txt")


# ===========================================================================
# TUNABLE MATCHING PATTERNS
# ---------------------------------------------------------------------------
# Backup spend is scattered across several Cost Explorer services and is only
# identifiable by USAGE_TYPE. Cost Explorer cannot substring-match usage types
# in a dimension filter, so the approach is: filter server-side to the services
# that can carry backup line items, then classify client-side here.
#
# The script DISCOVERS the account's real usage-type strings via
# get_dimension_values and matches them against these patterns, rather than
# assuming a hardcoded list stays current. Whatever matched (and what did not)
# is reported on the Notes sheet, so drift is visible instead of silent.
#
# Edit this block to tune the report. Nothing else should need changing.
# ===========================================================================

# Cost Explorer SERVICE dimension values that can contain backup charges.
# These strings are exact — "EC2 - Other" is where EBS snapshot spend lives,
# NOT "Amazon Elastic Compute Cloud".
BACKUP_SERVICES: List[str] = [
    "AWS Backup",
    "EC2 - Other",
    "Amazon Relational Database Service",
    "Amazon DynamoDB",
    "Amazon Elastic File System",
    "Amazon FSx",
    "Amazon Simple Storage Service",
    "Amazon Redshift",
]

# Report categories, in the order they appear on the Summary sheet.
CAT_VAULT = "Vault storage"
CAT_SNAPSHOTS = "Snapshots"
CAT_RDS = "RDS backup"
CAT_DDB = "DynamoDB backup"
CAT_RESTORE = "Restores and transfer"
CAT_OTHER = "Other backup"
CATEGORIES: List[str] = [
    CAT_VAULT,
    CAT_SNAPSHOTS,
    CAT_RDS,
    CAT_DDB,
    CAT_RESTORE,
    CAT_OTHER,
]

# (category, label, service filter or None for any, regex against USAGE_TYPE)
#
# ORDER MATTERS — first match wins. Two orderings are load-bearing:
#   1. "ChargedBackupUsage" (RDS) CONTAINS "BackupUsage" (DynamoDB/FSx) as a
#      substring. RDS must be tested first or every RDS backup line lands in
#      the DynamoDB bucket.
#   2. Restore and transfer patterns are tested before the generic storage
#      ones so a restore is never counted as storage.
USAGE_TYPE_RULES: List[Tuple[str, str, Optional[str], str]] = [
    # --- restores, copies and transfer (any service) ----------------------
    (CAT_RESTORE, "Restore", None, r"Restore"),
    (CAT_RESTORE, "Cross-region copy", None, r"CrossRegion"),
    (CAT_RESTORE, "Cross-account copy", None, r"CrossAccount"),
    (CAT_RESTORE, "Data transfer", None, r"DataTransfer|Data-Transfer|\bDTO\b"),
    (CAT_RESTORE, "Early delete", None, r"EarlyDelete"),

    # --- AWS Backup vault storage ----------------------------------------
    (CAT_VAULT, "Warm vault storage", "AWS Backup", r"WarmStorage"),
    (CAT_VAULT, "Cold vault storage", "AWS Backup", r"ColdStorage"),
    (CAT_VAULT, "Backup Audit Manager", "AWS Backup", r"AuditManager|BackupAudit"),
    # Catch-all for any other AWS Backup line item: every usage type inside
    # the AWS Backup service is backup spend by definition.
    (CAT_VAULT, "Other AWS Backup storage", "AWS Backup", r".*"),

    # --- EBS snapshots (billed under "EC2 - Other") -----------------------
    (CAT_SNAPSHOTS, "EBS snapshot storage", None, r"SnapshotUsage"),
    (CAT_SNAPSHOTS, "EBS snapshot archive", None, r"SnapshotArchiveStorage"),

    # --- RDS / Aurora ------  MUST precede the DynamoDB BackupUsage rule --
    (CAT_RDS, "RDS backup storage beyond free tier", None, r"ChargedBackupUsage"),
    (CAT_RDS, "RDS snapshot export to S3", None, r"SnapshotExport"),

    # --- DynamoDB ---------------------------------------------------------
    (CAT_DDB, "DynamoDB point-in-time recovery", None, r"PITR"),
    (CAT_DDB, "DynamoDB on-demand backup", None, r"BackupUsage|TimedBackupStorage"),

    # --- S3 archive storage classes --------------------------------------
    (CAT_OTHER, "S3 Glacier / Deep Archive", "Amazon Simple Storage Service",
     r"Glacier|DeepArchive"),

    # --- Redshift ---------------------------------------------------------
    (CAT_OTHER, "Redshift snapshot storage", "Amazon Redshift", r"Backup|Snapshot"),

    # --- EFS / FSx backup -------------------------------------------------
    (CAT_OTHER, "EFS backup storage", "Amazon Elastic File System", r"Backup"),
    (CAT_OTHER, "FSx backup storage", "Amazon FSx", r"Backup"),
]

# A snapshot with no surviving source volume and older than this is flagged.
ORPHAN_AGE_DAYS = 90

# Page caps. Cost Explorer bills $0.01 per request, and a very large estate
# can hold a lot of recovery points; these bound a run without failing it.
MAX_CE_PAGES = 40
MAX_RECOVERY_POINT_PAGES = 10
MAX_SNAPSHOT_PAGES = 20

# ===========================================================================
# End of tunable block.
# ===========================================================================


GIB = 1024 ** 3
BOTO_CONFIG = Config(retries={"max_attempts": 5, "mode": "standard"})

# Collected as the run proceeds and printed on the Notes sheet.
NOTES: List[Tuple[str, str]] = []


def note(topic: str, detail: str) -> None:
    """Record something the reader should know: a skip, a denial, a caveat."""
    NOTES.append((topic, detail))


def gib(num_bytes: Optional[float]) -> Optional[float]:
    if num_bytes is None:
        return None
    return round(num_bytes / GIB, 3)


def iso_date(d: dt.date) -> str:
    return d.strftime("%Y-%m-%d")


def months_back(now: dt.datetime, months: int) -> Tuple[str, str]:
    """Cost Explorer window: first day of the month `months - 1` back, through
    today (end is EXCLUSIVE, so data runs to yesterday). All UTC — Cost
    Explorer wants calendar dates, not times."""
    year = now.year
    month = now.month - (months - 1)
    while month <= 0:
        month += 12
        year -= 1
    return iso_date(dt.date(year, month, 1)), iso_date(now.date())


def classify(service: str, usage_type: str) -> Optional[Tuple[str, str]]:
    """Map a (SERVICE, USAGE_TYPE) pair to (category, rule label).

    Returns None when nothing matched, which means the line is ordinary
    non-backup spend inside a service we had to query wholesale.
    """
    for category, label, service_filter, pattern in USAGE_TYPE_RULES:
        if service_filter is not None and service != service_filter:
            continue
        if re.search(pattern, usage_type, re.IGNORECASE):
            return category, label
    return None


# ---------------------------------------------------------------------------
# Cost Explorer
# ---------------------------------------------------------------------------

class SpendData:
    def __init__(self) -> None:
        self.available = False
        self.error: Optional[str] = None
        # month -> category -> dollars
        self.by_month: Dict[str, Dict[str, float]] = defaultdict(
            lambda: defaultdict(float)
        )
        # (month, service, usage_type) -> dollars, for the detail sheet
        self.detail: List[Tuple[str, str, str, str, float]] = []
        # usage_type -> dollars, for the top-10 table
        self.by_usage_type: Dict[str, float] = defaultdict(float)
        self.matched_rules: Dict[str, float] = defaultdict(float)
        self.unmatched_usage_types: Dict[str, float] = defaultdict(float)
        self.discovered_usage_types: int = 0
        self.currency = "USD"

    @property
    def total(self) -> float:
        return sum(sum(c.values()) for c in self.by_month.values())

    def category_totals(self) -> Dict[str, float]:
        out: Dict[str, float] = {c: 0.0 for c in CATEGORIES}
        for cats in self.by_month.values():
            for cat, amount in cats.items():
                out[cat] = out.get(cat, 0.0) + amount
        return out


def explain_ce_error(err: Exception) -> str:
    raw = str(err)
    name = err.__class__.__name__
    if "AccessDenied" in raw or "AccessDenied" in name or "not authorized" in raw:
        return (
            "Cost Explorer access denied. In an AWS Organization the management "
            "account often disables Cost Explorer for member accounts, and the "
            "ce:GetCostAndUsage permission is separate from ReadOnlyAccess. "
            "Spend figures are unavailable; the inventory sheets are unaffected."
        )
    if "DataUnavailable" in raw:
        return (
            "Cost Explorer is enabled but has no data yet. A newly enabled "
            "account takes about 24 hours to populate. Spend figures are "
            "unavailable; the inventory sheets are unaffected."
        )
    if "OptInRequired" in raw or "not enabled" in raw.lower():
        return (
            "Cost Explorer is not enabled for this account. Enable it in the "
            "Billing console, wait about 24 hours, then re-run. The inventory "
            "sheets are unaffected."
        )
    return f"Cost Explorer query failed: {raw}. The inventory sheets are unaffected."


def _spend_filter(linked_account: Optional[str] = None) -> Dict[str, Any]:
    """Server-side Cost Explorer filter.

    Always narrows to the services that can carry backup line items. When a
    linked account is given, that is AND-ed on so the same query can be
    reused per account without touching the GroupBy dimensions.
    """
    service = {"Dimensions": {"Key": "SERVICE", "Values": BACKUP_SERVICES}}
    if linked_account is None:
        return service
    return {"And": [
        service,
        {"Dimensions": {"Key": "LINKED_ACCOUNT", "Values": [linked_account]}},
    ]}


def list_org_accounts(session: Any) -> Optional[List[Dict[str, str]]]:
    """Every ACTIVE account in the organisation, or None.

    None means this is not a management account, or the caller lacks
    organizations:ListAccounts. Either way the report falls back to the
    consolidated view, which is still org-wide in total — just not split
    per account.
    """
    try:
        org = session.client("organizations", region_name="us-east-1",
                             config=BOTO_CONFIG)
        accounts: List[Dict[str, str]] = []
        token = None
        while True:
            kwargs = {"NextToken": token} if token else {}
            resp = org.list_accounts(**kwargs)
            for a in resp.get("Accounts", []):
                if a.get("Status") != "ACTIVE":
                    continue
                accounts.append({"id": a.get("Id", ""),
                                 "name": a.get("Name", "")})
            token = resp.get("NextToken")
            if not token:
                break
        return sorted(accounts, key=lambda a: a["name"].lower()) or None
    except Exception as err:  # noqa: BLE001
        raw = str(err)
        if "AWSOrganizationsNotInUseException" in raw:
            note("Organization",
                 "This account is not part of an AWS Organization; the report "
                 "covers this account only.")
        elif "AccessDenied" in raw or "not authorized" in raw:
            note("Organization",
                 "Could not list organization accounts (access denied). Run "
                 "from the management account with organizations:ListAccounts "
                 "to break spend out per account. Totals from a management "
                 "account are still org-wide; from a member account they "
                 "cover that account only.")
        else:
            note("Organization",
                 f"Organization lookup failed ({err.__class__.__name__}); the "
                 "report covers whatever the credentials can see.")
        return None


def fetch_spend_by_account(session: Any, months: int, now: dt.datetime,
                           accounts: List[Dict[str, str]]
                           ) -> "Dict[str, SpendData]":
    """Re-run the spend query once per linked account.

    Each pass is one more Cost Explorer request at $0.01. Deliberately a
    loop rather than a LINKED_ACCOUNT GroupBy: the two GroupBy slots are
    already used by SERVICE and USAGE_TYPE, and dropping either would break
    the classification rules.
    """
    out: Dict[str, SpendData] = {}
    for account in accounts:
        account_id = account["id"]
        out[account_id] = fetch_spend(session, months, now,
                                      linked_account=account_id,
                                      discover=False)
    denied = [a["id"] for a in accounts if not out[a["id"]].available]
    if denied:
        note("Organization",
             f"{len(denied)} of {len(accounts)} accounts returned no Cost "
             "Explorer data; their rows show as unavailable and are excluded "
             "from the per-account table.")
    return out


def discover_usage_types(ce: Any, start: str, end: str) -> List[str]:
    """Ask Cost Explorer which USAGE_TYPE values this account actually has.

    Matching against discovered values beats a hardcoded list, which silently
    goes stale as AWS renames and adds usage types.
    """
    values: List[str] = []
    token: Optional[str] = None
    pages = 0
    while pages < MAX_CE_PAGES:
        kwargs: Dict[str, Any] = {
            "TimePeriod": {"Start": start, "End": end},
            "Dimension": "USAGE_TYPE",
        }
        if token:
            kwargs["NextPageToken"] = token
        resp = ce.get_dimension_values(**kwargs)
        values.extend(v.get("Value", "") for v in resp.get("DimensionValues", []))
        token = resp.get("NextPageToken")
        pages += 1
        if not token:
            break
    if token:
        note("Cost Explorer",
             f"Usage-type discovery stopped at the {MAX_CE_PAGES}-page cap; "
             "some usage types may not be listed on this sheet. Spend totals "
             "are unaffected — they come from a separate query.")
    return values


def fetch_spend(session: Any, months: int, now: dt.datetime,
                linked_account: Optional[str] = None,
                discover: bool = True) -> SpendData:
    """Backup spend for the whole payer scope, or for ONE linked account.

    Cost Explorer in a management account reports consolidated billing for
    every member account, so the unfiltered call is already org-wide. The
    per-account breakdown re-runs the SAME query with a LINKED_ACCOUNT
    filter added, which keeps classification byte-identical rather than
    regrouping — Cost Explorer allows only two GroupBy dimensions and both
    are already spent on SERVICE and USAGE_TYPE, which the rules need.
    """
    data = SpendData()
    start, end = months_back(now, months)
    # Cost Explorer is a global service anchored in us-east-1; one call covers
    # the whole account regardless of where the workloads run.
    ce = session.client("ce", region_name="us-east-1", config=BOTO_CONFIG)

    if not discover:
        # Per-account passes skip discovery: the usage types are the same
        # ones the org-wide pass already listed, and each call costs $0.01.
        pass
    else:
      try:
        discovered = discover_usage_types(ce, start, end)
        data.discovered_usage_types = len(discovered)
        matched_preview = [u for u in discovered if classify("AWS Backup", u) or
                           classify("EC2 - Other", u)]
        note("Cost Explorer",
             f"Discovered {len(discovered)} usage types in the window; "
             f"{len(matched_preview)} matched a backup pattern on a first pass. "
             "Final attribution is per line item and is listed below.")
      except Exception as err:  # noqa: BLE001 - discovery is best effort
        note("Cost Explorer",
             f"Usage-type discovery failed ({err.__class__.__name__}); "
             "classification still ran against the live line items.")

    try:
        token: Optional[str] = None
        pages = 0
        while pages < MAX_CE_PAGES:
            kwargs: Dict[str, Any] = {
                "TimePeriod": {"Start": start, "End": end},
                "Granularity": "MONTHLY",
                "Metrics": ["UnblendedCost"],
                "Filter": _spend_filter(linked_account),
                "GroupBy": [
                    {"Type": "DIMENSION", "Key": "SERVICE"},
                    {"Type": "DIMENSION", "Key": "USAGE_TYPE"},
                ],
            }
            if token:
                kwargs["NextPageToken"] = token
            resp = ce.get_cost_and_usage(**kwargs)

            for result in resp.get("ResultsByTime", []):
                month = result.get("TimePeriod", {}).get("Start", "")[:7]
                for group in result.get("Groups", []):
                    keys = group.get("Keys", ["", ""])
                    service = keys[0] if keys else ""
                    usage_type = keys[1] if len(keys) > 1 else ""
                    metric = group.get("Metrics", {}).get("UnblendedCost", {})
                    data.currency = metric.get("Unit", data.currency) or data.currency
                    try:
                        amount = float(metric.get("Amount", "0"))
                    except (TypeError, ValueError):
                        continue
                    if amount == 0:
                        continue

                    hit = classify(service, usage_type)
                    if hit is None:
                        # Ordinary non-backup spend inside a service we had to
                        # query wholesale (e.g. plain S3 storage). Recorded so
                        # the Notes sheet can show what was set aside.
                        data.unmatched_usage_types[f"{service} / {usage_type}"] += amount
                        continue
                    category, label = hit
                    data.by_month[month][category] += amount
                    data.by_usage_type[f"{service} / {usage_type}"] += amount
                    data.matched_rules[label] += amount
                    data.detail.append((month, service, usage_type, category, amount))

            token = resp.get("NextPageToken")
            pages += 1
            if not token:
                break

        data.available = True
        if linked_account is None:
            note("Cost Explorer",
                 f"Window {start} to {end} (end exclusive), MONTHLY granularity, "
                 f"UnblendedCost, {pages} request(s) at $0.01 each.")
    except Exception as err:  # noqa: BLE001 - never fail the run on spend
        data.error = explain_ce_error(err)
        if linked_account is None:
            note("Cost Explorer", data.error)

    return data


# ---------------------------------------------------------------------------
# Inventory
# ---------------------------------------------------------------------------

def assume_account_session(session: Any, account_id: str,
                           role_name: str) -> Optional[Any]:
    """A boto3 session holding temporary credentials in a member account.

    Because the customer runs this script themselves, the role only has to
    trust their OWN management account — no external ID and no third-party
    trust, which removes the whole class of "wrong ExternalId" failures
    that third-party assume-role integrations are prone to.

    Returns None on failure, with the reason recorded. A single
    inaccessible account must never fail the run.
    """
    role_arn = "arn:aws:iam::%s:role/%s" % (account_id, role_name)
    try:
        sts = session.client("sts", config=BOTO_CONFIG)
        creds = sts.assume_role(
            RoleArn=role_arn,
            RoleSessionName="backup-cost-report",
        )["Credentials"]
        return boto3.Session(
            aws_access_key_id=creds["AccessKeyId"],
            aws_secret_access_key=creds["SecretAccessKey"],
            aws_session_token=creds["SessionToken"],
        )
    except Exception as err:  # noqa: BLE001
        raw = str(err)
        if "AccessDenied" in raw or "not authorized" in raw:
            detail = ("access denied assuming %s. Check the role exists in "
                      "that account and trusts this one. IAM role names are "
                      "CASE-SENSITIVE inside an ARN." % role_arn)
        else:
            detail = "%s assuming %s" % (err.__class__.__name__, role_arn)
        note("Organization inventory", "%s: %s" % (account_id, detail))
        return None


def scan_account_inventory(session: Any, regions: Sequence[str],
                           now: dt.datetime, account_id: str
                           ) -> Dict[str, List[Dict[str, Any]]]:
    """Run every inventory scanner against one account and tag the rows.

    The scanners all take a session as their first argument, so covering a
    whole organisation is a loop over sessions rather than a rewrite.
    """
    result = {
        "vaults": scan_vaults(session, regions),
        "snapshots": scan_ebs_snapshots(session, regions, now),
        "rds": scan_rds(session, regions),
        "ddb": scan_dynamodb(session, regions),
        "redshift": scan_redshift(session, regions),
    }
    for rows in result.values():
        for row in rows:
            row["account"] = account_id
    return result


def enabled_regions(session: Any) -> List[str]:
    """Regions this account has enabled, including opt-in regions.

    Enumerating properly matters: a hand-picked region list is how backup
    storage in a region nobody remembers goes unnoticed, which is exactly what
    this report exists to surface.
    """
    try:
        ec2 = session.client("ec2", region_name="us-east-1", config=BOTO_CONFIG)
        resp = ec2.describe_regions(AllRegions=False)
        regions = sorted(r["RegionName"] for r in resp.get("Regions", []))
        if regions:
            return regions
    except Exception as err:  # noqa: BLE001
        note("Regions",
             f"Could not list enabled regions ({err.__class__.__name__}); "
             "fell back to the session's region only. Anything outside it is "
             "NOT in this report.")
    fallback = session.region_name or "us-east-1"
    return [fallback]


def scan_vaults(session: Any, regions: Sequence[str]) -> List[Dict[str, Any]]:
    """AWS Backup vaults with recovery-point counts, size by tier, and age."""
    rows: List[Dict[str, Any]] = []
    for region in regions:
        try:
            client = session.client("backup", region_name=region, config=BOTO_CONFIG)
            vaults: List[Dict[str, Any]] = []
            token = None
            while True:
                kwargs = {"NextToken": token} if token else {}
                resp = client.list_backup_vaults(**kwargs)
                vaults.extend(resp.get("BackupVaultList", []))
                token = resp.get("NextToken")
                if not token:
                    break
        except Exception as err:  # noqa: BLE001
            note("AWS Backup",
                 f"{region}: vault listing skipped ({err.__class__.__name__}).")
            continue

        for vault in vaults:
            name = vault.get("BackupVaultName", "")
            row: Dict[str, Any] = {
                "region": region,
                "vault": name,
                "recovery_points": vault.get("NumberOfRecoveryPoints", 0),
                "warm_gib": 0.0,
                "cold_gib": 0.0,
                "total_gib": 0.0,
                "oldest": None,
                "newest": None,
                "locked": "yes" if vault.get("Locked") else "no",
                "truncated": "no",
            }
            try:
                token = None
                pages = 0
                while pages < MAX_RECOVERY_POINT_PAGES:
                    kwargs: Dict[str, Any] = {
                        "BackupVaultName": name,
                        "MaxResults": 1000,
                    }
                    if token:
                        kwargs["NextToken"] = token
                    resp = client.list_recovery_points_by_backup_vault(**kwargs)
                    for rp in resp.get("RecoveryPoints", []):
                        size = rp.get("BackupSizeInBytes") or 0
                        storage_class = (rp.get("StorageClass") or "WARM").upper()
                        if storage_class == "COLD":
                            row["cold_gib"] += size / GIB
                        elif storage_class != "DELETED":
                            row["warm_gib"] += size / GIB
                        created = rp.get("CreationDate")
                        if created:
                            if row["oldest"] is None or created < row["oldest"]:
                                row["oldest"] = created
                            if row["newest"] is None or created > row["newest"]:
                                row["newest"] = created
                    token = resp.get("NextToken")
                    pages += 1
                    if not token:
                        break
                if token:
                    row["truncated"] = "yes"
                    note("AWS Backup",
                         f"{region}/{name}: stopped at {MAX_RECOVERY_POINT_PAGES} "
                         "pages of recovery points; sizes for this vault are a "
                         "lower bound.")
            except Exception as err:  # noqa: BLE001
                note("AWS Backup",
                     f"{region}/{name}: recovery points unreadable "
                     f"({err.__class__.__name__}); size left blank.")
            row["warm_gib"] = round(row["warm_gib"], 3)
            row["cold_gib"] = round(row["cold_gib"], 3)
            row["total_gib"] = round(row["warm_gib"] + row["cold_gib"], 3)
            rows.append(row)
    return rows


def scan_ebs_snapshots(session: Any, regions: Sequence[str],
                       now: dt.datetime) -> List[Dict[str, Any]]:
    """Every self-owned EBS snapshot, with age and whether its volume survives."""
    rows: List[Dict[str, Any]] = []
    for region in regions:
        try:
            ec2 = session.client("ec2", region_name=region, config=BOTO_CONFIG)
        except Exception:  # noqa: BLE001
            continue

        live_volumes = set()
        try:
            token = None
            while True:
                kwargs: Dict[str, Any] = {"MaxResults": 500}
                if token:
                    kwargs["NextToken"] = token
                resp = ec2.describe_volumes(**kwargs)
                for v in resp.get("Volumes", []):
                    live_volumes.add(v.get("VolumeId"))
                token = resp.get("NextToken")
                if not token:
                    break
        except Exception as err:  # noqa: BLE001
            note("EBS",
                 f"{region}: volumes unreadable ({err.__class__.__name__}); "
                 "'source volume exists' is reported as unknown here.")
            live_volumes = set()
            volumes_known = False
        else:
            volumes_known = True

        try:
            token = None
            pages = 0
            while pages < MAX_SNAPSHOT_PAGES:
                # OwnerIds=self is essential. Without it this returns every
                # public and shared snapshot in the region.
                kwargs: Dict[str, Any] = {"OwnerIds": ["self"], "MaxResults": 1000}
                if token:
                    kwargs["NextToken"] = token
                resp = ec2.describe_snapshots(**kwargs)
                for snap in resp.get("Snapshots", []):
                    started = snap.get("StartTime")
                    age_days = None
                    if started:
                        # boto3 returns aware datetimes; normalise anything
                        # naive to UTC so the subtraction can never raise.
                        if started.tzinfo is None:
                            started = started.replace(tzinfo=dt.timezone.utc)
                        age_days = max(0, (now - started).days)
                    volume_id = snap.get("VolumeId") or ""
                    tags = {t.get("Key", ""): t.get("Value", "")
                            for t in snap.get("Tags", [])}
                    # AWS Backup tags what it creates; older/other paths only
                    # set the description. Check both.
                    in_backup = (
                        "aws:backup:source-resource" in tags
                        or "AWS Backup" in (snap.get("Description") or "")
                    )
                    if not volumes_known:
                        exists = "unknown"
                    elif volume_id and volume_id in live_volumes:
                        exists = "yes"
                    else:
                        exists = "no"
                    rows.append({
                        "region": region,
                        "snapshot": snap.get("SnapshotId", ""),
                        "volume": volume_id,
                        "size_gib": snap.get("VolumeSize", 0),
                        "created": started,
                        "age_days": age_days,
                        "volume_exists": exists,
                        "in_aws_backup": "yes" if in_backup else "no",
                        "description": (snap.get("Description") or "")[:120],
                    })
                token = resp.get("NextToken")
                pages += 1
                if not token:
                    break
            if token:
                note("EBS",
                     f"{region}: stopped at {MAX_SNAPSHOT_PAGES} pages of "
                     "snapshots; the list for this region is partial.")
        except Exception as err:  # noqa: BLE001
            note("EBS",
                 f"{region}: snapshots unreadable ({err.__class__.__name__}).")
    return rows


def is_orphan(row: Dict[str, Any]) -> bool:
    return (row.get("volume_exists") == "no"
            and (row.get("age_days") or 0) > ORPHAN_AGE_DAYS)


def scan_rds(session: Any, regions: Sequence[str]) -> List[Dict[str, Any]]:
    """RDS/Aurora retention settings and manual snapshot inventory."""
    rows: List[Dict[str, Any]] = []
    for region in regions:
        try:
            rds = session.client("rds", region_name=region, config=BOTO_CONFIG)
        except Exception:  # noqa: BLE001
            continue

        # Clusters first, so their member instances can be skipped below —
        # describe_db_instances returns Aurora members too and they would
        # otherwise be counted twice.
        members = set()
        try:
            marker = None
            while True:
                kwargs = {"Marker": marker} if marker else {}
                resp = rds.describe_db_clusters(**kwargs)
                for c in resp.get("DBClusters", []):
                    for m in c.get("DBClusterMembers", []):
                        members.add(m.get("DBInstanceIdentifier"))
                    rows.append({
                        "region": region,
                        "kind": "cluster",
                        "identifier": c.get("DBClusterIdentifier", ""),
                        "engine": c.get("Engine", ""),
                        "retention_days": c.get("BackupRetentionPeriod", 0),
                        "manual_snapshots": 0,
                        "deletion_protection": "yes" if c.get("DeletionProtection") else "no",
                    })
                marker = resp.get("Marker")
                if not marker:
                    break
        except Exception as err:  # noqa: BLE001
            note("RDS", f"{region}: clusters unreadable ({err.__class__.__name__}).")

        try:
            marker = None
            while True:
                kwargs = {"Marker": marker} if marker else {}
                resp = rds.describe_db_instances(**kwargs)
                for i in resp.get("DBInstances", []):
                    ident = i.get("DBInstanceIdentifier", "")
                    if ident in members:
                        continue
                    rows.append({
                        "region": region,
                        "kind": "instance",
                        "identifier": ident,
                        "engine": i.get("Engine", ""),
                        "retention_days": i.get("BackupRetentionPeriod", 0),
                        "manual_snapshots": 0,
                        "deletion_protection": "yes" if i.get("DeletionProtection") else "no",
                    })
                marker = resp.get("Marker")
                if not marker:
                    break
        except Exception as err:  # noqa: BLE001
            note("RDS", f"{region}: instances unreadable ({err.__class__.__name__}).")

        # Manual snapshots, counted against their source.
        counts: Dict[str, int] = defaultdict(int)
        for call, key, id_key in (
            (rds.describe_db_snapshots, "DBSnapshots", "DBInstanceIdentifier"),
            (rds.describe_db_cluster_snapshots, "DBClusterSnapshots", "DBClusterIdentifier"),
        ):
            try:
                marker = None
                while True:
                    kwargs: Dict[str, Any] = {"SnapshotType": "manual"}
                    if marker:
                        kwargs["Marker"] = marker
                    resp = call(**kwargs)
                    for s in resp.get(key, []):
                        counts[s.get(id_key, "")] += 1
                    marker = resp.get("Marker")
                    if not marker:
                        break
            except Exception as err:  # noqa: BLE001
                note("RDS",
                     f"{region}: manual snapshots unreadable "
                     f"({err.__class__.__name__}).")
        for row in rows:
            if row["region"] == region:
                row["manual_snapshots"] = counts.get(row["identifier"], 0)
    return rows


def scan_dynamodb(session: Any, regions: Sequence[str]) -> List[Dict[str, Any]]:
    """DynamoDB tables with PITR status and on-demand backup counts."""
    rows: List[Dict[str, Any]] = []
    for region in regions:
        try:
            ddb = session.client("dynamodb", region_name=region, config=BOTO_CONFIG)
            tables: List[str] = []
            start = None
            while True:
                kwargs = {"ExclusiveStartTableName": start} if start else {}
                resp = ddb.list_tables(**kwargs)
                tables.extend(resp.get("TableNames", []))
                start = resp.get("LastEvaluatedTableName")
                if not start:
                    break
        except Exception as err:  # noqa: BLE001
            note("DynamoDB",
                 f"{region}: tables unreadable ({err.__class__.__name__}).")
            continue

        for table in tables:
            size_bytes = None
            try:
                desc = ddb.describe_table(TableName=table).get("Table", {})
                size_bytes = desc.get("TableSizeBytes")
            except Exception:  # noqa: BLE001
                pass  # table may be CREATING or have just been deleted

            pitr = "unknown"
            try:
                cb = ddb.describe_continuous_backups(TableName=table)
                status = (cb.get("ContinuousBackupsDescription", {})
                            .get("PointInTimeRecoveryDescription", {})
                            .get("PointInTimeRecoveryStatus"))
                pitr = "enabled" if status == "ENABLED" else "disabled"
            except Exception:  # noqa: BLE001
                pass  # some account types disallow this call

            backups = 0
            try:
                start_arn = None
                while True:
                    kwargs: Dict[str, Any] = {"TableName": table}
                    if start_arn:
                        kwargs["ExclusiveStartBackupArn"] = start_arn
                    resp = ddb.list_backups(**kwargs)
                    backups += len(resp.get("BackupSummaries", []))
                    start_arn = resp.get("LastEvaluatedBackupArn")
                    if not start_arn:
                        break
            except Exception:  # noqa: BLE001
                backups = -1  # rendered as "unknown"

            rows.append({
                "region": region,
                "table": table,
                "size_gib": gib(size_bytes),
                "pitr": pitr,
                "on_demand_backups": backups,
            })
    return rows


def scan_redshift(session: Any, regions: Sequence[str]) -> List[Dict[str, Any]]:
    """Redshift manual snapshot inventory."""
    rows: List[Dict[str, Any]] = []
    for region in regions:
        try:
            rs = session.client("redshift", region_name=region, config=BOTO_CONFIG)
            marker = None
            counts: Dict[str, int] = defaultdict(int)
            while True:
                kwargs: Dict[str, Any] = {"SnapshotType": "manual"}
                if marker:
                    kwargs["Marker"] = marker
                resp = rs.describe_cluster_snapshots(**kwargs)
                for s in resp.get("Snapshots", []):
                    counts[s.get("ClusterIdentifier", "")] += 1
                marker = resp.get("Marker")
                if not marker:
                    break
            for cluster, count in sorted(counts.items()):
                rows.append({
                    "region": region,
                    "cluster": cluster,
                    "manual_snapshots": count,
                })
        except Exception as err:  # noqa: BLE001
            name = err.__class__.__name__
            # Redshift is not enabled in every region; that is not an error.
            if "AccessDenied" in str(err) or "UnauthorizedOperation" in str(err):
                note("Redshift", f"{region}: access denied ({name}).")
    return rows

# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="44546A")
TITLE_FONT = Font(bold=True, size=14)
WARN_FILL = PatternFill("solid", fgColor="FCE4E4")
MONEY = '#,##0.00'
SIZE_FMT = '#,##0.000'


def write_header(ws: Any, row: int, headers: Sequence[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def finish_sheet(ws: Any, widths: Sequence[float], account_id: str,
                 header_row: int = 1, freeze: Optional[str] = None) -> None:
    """Column widths, freeze panes and PRINT SETUP.

    The workbook is meant to be exported to PDF from Excel (File > Export), so
    every sheet gets a print area, landscape orientation, fit-to-width, a
    repeating header row and a footer. Without these a wide sheet prints
    across several unreadable pages.
    """
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    if freeze:
        ws.freeze_panes = freeze

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "%d:%d" % (header_row, header_row)

    last_col = get_column_letter(max(1, ws.max_column))
    ws.print_area = "A1:%s%d" % (last_col, max(1, ws.max_row))

    ws.oddFooter.left.text = "AWS account %s" % account_id
    ws.oddFooter.right.text = "Generated &D  |  Page &P of &N"


def _naive(value: Any) -> Any:
    """Excel cannot store a timezone-aware datetime."""
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)
    return value


def sheet_summary(wb: Any, spend: SpendData, snapshots: List[Dict[str, Any]],
                  account_id: str, start: str, end: str, months: int,
                  org: Optional[Dict[str, Any]] = None) -> None:
    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = "AWS backup spend"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Account %s   |   %s to %s (end exclusive)   |   %d months" % (
        account_id, start, end, months)

    row = 4
    if not spend.available:
        ws.cell(row=row, column=1, value="Cost Explorer data unavailable").font = Font(bold=True)
        ws.cell(row=row + 1, column=1,
                value=spend.error or "Cost Explorer returned no data.")
        ws.cell(row=row + 2, column=1,
                value="The inventory sheets in this workbook were still produced "
                      "and are unaffected.")
        row += 4
    else:
        ws.cell(row=row, column=1, value="Total backup spend in window").font = Font(bold=True)
        total_cell = ws.cell(row=row, column=2, value=round(spend.total, 2))
        total_cell.number_format = MONEY
        total_cell.font = Font(bold=True, size=12)
        ws.cell(row=row, column=3, value=spend.currency)
        row += 1
        monthly_avg = spend.total / max(1, len(spend.by_month))
        ws.cell(row=row, column=1, value="Monthly average")
        ws.cell(row=row, column=2, value=round(monthly_avg, 2)).number_format = MONEY
        row += 2

        if org:
            ws.cell(row=row, column=1, value="By account").font = Font(bold=True)
            row += 1
            write_header(ws, row, ["Account", "Name", "Cost", "% of total"])
            row += 1
            for entry in org["rows"]:
                ws.cell(row=row, column=1, value=entry["id"])
                ws.cell(row=row, column=2, value=entry["name"])
                ws.cell(row=row, column=3,
                        value=round(entry["total"], 2)).number_format = MONEY
                ws.cell(row=row, column=4,
                        value=round(entry["total"] / (spend.total or 1.0), 4)
                        ).number_format = '0.0%'
                row += 1
            row += 1

        ws.cell(row=row, column=1, value="By category").font = Font(bold=True)
        row += 1
        write_header(ws, row, ["Category", "Cost", "% of total"])
        row += 1
        totals = spend.category_totals()
        grand = spend.total or 1.0
        for category in CATEGORIES:
            amount = totals.get(category, 0.0)
            if amount == 0:
                continue
            ws.cell(row=row, column=1, value=category)
            ws.cell(row=row, column=2, value=round(amount, 2)).number_format = MONEY
            ws.cell(row=row, column=3, value=round(amount / grand, 4)).number_format = '0.0%'
            row += 1
        row += 1

        ws.cell(row=row, column=1, value="Monthly trend").font = Font(bold=True)
        row += 1
        write_header(ws, row, ["Month"] + CATEGORIES + ["Total"])
        row += 1
        for month in sorted(spend.by_month):
            cats = spend.by_month[month]
            ws.cell(row=row, column=1, value=month)
            for i, category in enumerate(CATEGORIES, start=2):
                cell = ws.cell(row=row, column=i, value=round(cats.get(category, 0.0), 2))
                cell.number_format = MONEY
            total_cell = ws.cell(row=row, column=len(CATEGORIES) + 2,
                                 value=round(sum(cats.values()), 2))
            total_cell.number_format = MONEY
            total_cell.font = Font(bold=True)
            row += 1
        row += 1

        ws.cell(row=row, column=1, value="Top 10 usage types by cost").font = Font(bold=True)
        row += 1
        write_header(ws, row, ["Service / usage type", "Cost"])
        row += 1
        top = sorted(spend.by_usage_type.items(), key=lambda kv: kv[1], reverse=True)[:10]
        for name, amount in top:
            ws.cell(row=row, column=1, value=name)
            ws.cell(row=row, column=2, value=round(amount, 2)).number_format = MONEY
            row += 1
        row += 1

    orphans = [s for s in snapshots if is_orphan(s)]
    ws.cell(row=row, column=1, value="Orphaned EBS snapshots").font = Font(bold=True)
    row += 1
    ws.cell(row=row, column=1,
            value="Source volume deleted and older than %d days" % ORPHAN_AGE_DAYS)
    ws.cell(row=row, column=2, value=len(orphans))
    row += 1
    ws.cell(row=row, column=1, value="Total size (GiB)")
    ws.cell(row=row, column=2,
            value=round(sum(s.get("size_gib") or 0 for s in orphans), 3)).number_format = SIZE_FMT

    finish_sheet(ws, [46, 16, 14, 14, 14, 14, 14, 14], account_id, header_row=1)


def sheet_by_account(wb: Any, org: Dict[str, Any], account_id: str) -> None:
    """One row per organisation account, categories across.

    Only produced when the run could list the organisation, i.e. from a
    management account with organizations:ListAccounts.
    """
    ws = wb.create_sheet("By account")
    write_header(ws, 1, ["Account", "Name"] + CATEGORIES + ["Total"])
    row = 2
    for entry in org["rows"]:
        ws.cell(row=row, column=1, value=entry["id"])
        ws.cell(row=row, column=2, value=entry["name"])
        for i, category in enumerate(CATEGORIES, start=3):
            cell = ws.cell(row=row, column=i,
                           value=round(entry["categories"].get(category, 0.0), 2))
            cell.number_format = MONEY
        total = ws.cell(row=row, column=len(CATEGORIES) + 3,
                        value=round(entry["total"], 2))
        total.number_format = MONEY
        total.font = Font(bold=True)
        row += 1

    if org["unavailable"]:
        row += 1
        ws.cell(row=row, column=1,
                value="No Cost Explorer data returned for these accounts:"
                ).font = Font(bold=True)
        row += 1
        for entry in org["unavailable"]:
            ws.cell(row=row, column=1, value=entry["id"])
            ws.cell(row=row, column=2, value=entry["name"])
            row += 1

    widths = [16, 30] + [15] * len(CATEGORIES) + [15]
    finish_sheet(ws, widths, account_id, freeze="A2")


def sheet_monthly_detail(wb: Any, spend: SpendData, account_id: str) -> None:
    ws = wb.create_sheet("Monthly detail")
    write_header(ws, 1, ["Month", "Service", "Usage type", "Category", "Cost"])
    row = 2
    for month, service, usage_type, category, amount in sorted(spend.detail):
        ws.cell(row=row, column=1, value=month)
        ws.cell(row=row, column=2, value=service)
        ws.cell(row=row, column=3, value=usage_type)
        ws.cell(row=row, column=4, value=category)
        ws.cell(row=row, column=5, value=round(amount, 2)).number_format = MONEY
        row += 1
    if row == 2:
        ws.cell(row=2, column=1, value="No backup line items matched in this window.")
    finish_sheet(ws, [12, 34, 40, 22, 14], account_id, freeze="A2")


def sheet_vaults(wb: Any, vaults: List[Dict[str, Any]], account_id: str) -> None:
    ws = wb.create_sheet("Vaults")
    write_header(ws, 1, ["Account", "Region", "Vault", "Recovery points",
                         "Warm (GiB)", "Cold (GiB)", "Total (GiB)", "Oldest",
                         "Newest", "Locked", "Truncated"])
    row = 2
    for v in sorted(vaults, key=lambda r: (r.get("account", ""), r["region"],
                                           r["vault"])):
        ws.cell(row=row, column=1, value=v.get("account", ""))
        ws.cell(row=row, column=2, value=v["region"])
        ws.cell(row=row, column=3, value=v["vault"])
        ws.cell(row=row, column=4, value=v["recovery_points"])
        ws.cell(row=row, column=5, value=v["warm_gib"]).number_format = SIZE_FMT
        ws.cell(row=row, column=6, value=v["cold_gib"]).number_format = SIZE_FMT
        ws.cell(row=row, column=7, value=v["total_gib"]).number_format = SIZE_FMT
        for col, key in ((8, "oldest"), (9, "newest")):
            cell = ws.cell(row=row, column=col, value=_naive(v.get(key)))
            cell.number_format = 'yyyy-mm-dd'
        ws.cell(row=row, column=10, value=v["locked"])
        ws.cell(row=row, column=11, value=v["truncated"])
        row += 1
    if row == 2:
        ws.cell(row=2, column=1, value="No AWS Backup vaults found in the scanned regions.")
    finish_sheet(ws, [16, 16, 34, 16, 14, 14, 14, 14, 14, 10, 12], account_id,
                 freeze="B2")


def sheet_snapshots(wb: Any, snapshots: List[Dict[str, Any]], account_id: str) -> None:
    ws = wb.create_sheet("Snapshots")
    write_header(ws, 1, ["Account", "Region", "Snapshot", "Source volume",
                         "Size (GiB)", "Created", "Age (days)",
                         "Volume exists", "In AWS Backup", "Orphaned",
                         "Description"])
    row = 2
    for s in sorted(snapshots, key=lambda r: (r.get("account", ""), r["region"],
                                              r["snapshot"])):
        orphan = is_orphan(s)
        ws.cell(row=row, column=1, value=s.get("account", ""))
        ws.cell(row=row, column=2, value=s["region"])
        ws.cell(row=row, column=3, value=s["snapshot"])
        ws.cell(row=row, column=4, value=s["volume"])
        ws.cell(row=row, column=5, value=s["size_gib"]).number_format = '#,##0'
        cell = ws.cell(row=row, column=6, value=_naive(s.get("created")))
        cell.number_format = 'yyyy-mm-dd'
        ws.cell(row=row, column=7, value=s.get("age_days"))
        ws.cell(row=row, column=8, value=s["volume_exists"])
        ws.cell(row=row, column=9, value=s["in_aws_backup"])
        ws.cell(row=row, column=10, value="yes" if orphan else "no")
        ws.cell(row=row, column=11, value=s["description"])
        if orphan:
            # Highlight the whole row: these are the ones worth acting on.
            for col in range(1, 12):
                ws.cell(row=row, column=col).fill = WARN_FILL
        row += 1
    if row == 2:
        ws.cell(row=2, column=1, value="No self-owned EBS snapshots found.")
    finish_sheet(ws, [16, 16, 24, 22, 12, 14, 12, 14, 14, 12, 46], account_id,
                 freeze="B2")


def sheet_rds_dynamodb(wb: Any, rds: List[Dict[str, Any]],
                       ddb: List[Dict[str, Any]], redshift: List[Dict[str, Any]],
                       account_id: str) -> None:
    ws = wb.create_sheet("RDS and DynamoDB")

    ws["A1"] = "RDS / Aurora"
    ws["A1"].font = TITLE_FONT
    write_header(ws, 2, ["Account", "Region", "Type", "Identifier", "Engine",
                         "Automated retention (days)", "Manual snapshots",
                         "Deletion protection"])
    row = 3
    for r in sorted(rds, key=lambda x: (x.get("account", ""), x["region"],
                                        x["identifier"])):
        ws.cell(row=row, column=1, value=r.get("account", ""))
        ws.cell(row=row, column=2, value=r["region"])
        ws.cell(row=row, column=3, value=r["kind"])
        ws.cell(row=row, column=4, value=r["identifier"])
        ws.cell(row=row, column=5, value=r["engine"])
        retention = ws.cell(row=row, column=6, value=r["retention_days"])
        if not r["retention_days"]:
            # Retention 0 means automated backups are switched OFF.
            retention.fill = WARN_FILL
        ws.cell(row=row, column=7, value=r["manual_snapshots"])
        ws.cell(row=row, column=8, value=r["deletion_protection"])
        row += 1
    if row == 3:
        ws.cell(row=row, column=1, value="No RDS instances or clusters found.")
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="DynamoDB").font = TITLE_FONT
    row += 1
    write_header(ws, row, ["Account", "Region", "Table", "Size (GiB)", "PITR",
                           "On-demand backups"])
    row += 1
    for t in sorted(ddb, key=lambda x: (x.get("account", ""), x["region"],
                                        x["table"])):
        ws.cell(row=row, column=1, value=t.get("account", ""))
        ws.cell(row=row, column=2, value=t["region"])
        ws.cell(row=row, column=3, value=t["table"])
        ws.cell(row=row, column=4, value=t["size_gib"]).number_format = SIZE_FMT
        pitr = ws.cell(row=row, column=5, value=t["pitr"])
        if t["pitr"] == "disabled":
            pitr.fill = WARN_FILL
        count = t["on_demand_backups"]
        ws.cell(row=row, column=6,
                value="unknown" if count is not None and count < 0 else count)
        row += 1
    if not ddb:
        ws.cell(row=row, column=1, value="No DynamoDB tables found.")
        row += 1

    if redshift:
        row += 2
        ws.cell(row=row, column=1, value="Redshift").font = TITLE_FONT
        row += 1
        write_header(ws, row, ["Account", "Region", "Cluster",
                               "Manual snapshots"])
        row += 1
        for c in redshift:
            ws.cell(row=row, column=1, value=c.get("account", ""))
            ws.cell(row=row, column=2, value=c["region"])
            ws.cell(row=row, column=3, value=c["cluster"])
            ws.cell(row=row, column=4, value=c["manual_snapshots"])
            row += 1

    finish_sheet(ws, [16, 16, 14, 40, 20, 24, 18, 20], account_id, header_row=2)


def sheet_notes(wb: Any, spend: SpendData, account_id: str, regions: Sequence[str],
                start: str, end: str, generated: dt.datetime,
                org: Optional[Dict[str, Any]] = None,
                inventory_accounts: Optional[Sequence[str]] = None) -> None:
    ws = wb.create_sheet("Notes")
    if org:
        scope_label = ("Organization-wide, broken out across %d accounts "
                       "(consolidated billing)." % len(org["rows"]))
    else:
        scope_label = ("Whatever this account's Cost Explorer covers. From a "
                       "management account that is the whole organization; "
                       "from a member account it is this account only.")
    covered = list(inventory_accounts or [account_id])
    if len(covered) > 1:
        inventory_scope = ("%d accounts: %s. Gathered by assuming a read-only "
                           "role in each." % (len(covered), ", ".join(covered)))
    else:
        inventory_scope = ("This account only (%s). Vaults, snapshots, RDS and "
                           "DynamoDB are per-account APIs with no consolidated "
                           "view; pass --assume-role to cover the organization."
                           % covered[0])
    row = 1
    ws.cell(row=row, column=1, value="Run details").font = TITLE_FONT
    row += 2
    for label, value in (
        ("AWS account", account_id),
        ("Window", "%s to %s (end exclusive)" % (start, end)),
        ("Generated (UTC)", generated.strftime("%Y-%m-%d %H:%M:%S")),
        ("Regions scanned", "%d: %s" % (len(regions), ", ".join(regions))),
        ("Cost Explorer", "available" if spend.available else "UNAVAILABLE"),
        ("Spend scope", scope_label),
        ("Inventory scope", inventory_scope),
    ):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value).alignment = Alignment(wrap_text=True)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="How to read this report").font = TITLE_FONT
    row += 1
    for line in (
        "Money columns are real billed dollars from Cost Explorer (UnblendedCost). "
        "Nothing in this workbook is an estimate or a quote.",
        "Inventory sheets report counts, sizes and ages only. They are not priced - "
        "they exist to explain the bill and to surface storage that has been forgotten.",
        "S3 versioning or replication used as a backup is indistinguishable from "
        "primary S3 storage on the AWS bill, so it cannot appear here. Only the "
        "Glacier and Deep Archive storage classes are identifiable.",
        "RDS backup storage within the free allocation never appears on the bill, so "
        "an RDS instance with backups enabled can legitimately show no cost.",
        "A snapshot is flagged as orphaned when its source volume no longer exists "
        "and it is older than %d days." % ORPHAN_AGE_DAYS,
    ):
        cell = ws.cell(row=row, column=1, value="- " + line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.row_dimensions[row].height = 30
        row += 1

    if spend.matched_rules:
        row += 1
        ws.cell(row=row, column=1, value="Usage-type patterns that matched").font = TITLE_FONT
        row += 1
        write_header(ws, row, ["Pattern", "Cost"])
        row += 1
        for label, amount in sorted(spend.matched_rules.items(),
                                    key=lambda kv: kv[1], reverse=True):
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=round(amount, 2)).number_format = MONEY
            row += 1

    if spend.unmatched_usage_types:
        row += 1
        ws.cell(row=row, column=1,
                value="Set aside as non-backup spend").font = TITLE_FONT
        row += 1
        ws.cell(row=row, column=1,
                value="These line items sit inside a service that had to be queried "
                      "wholesale but matched no backup pattern. Listed so the "
                      "patterns can be tuned if something is missing.")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        row += 1
        write_header(ws, row, ["Service / usage type", "Cost (excluded)"])
        row += 1
        top = sorted(spend.unmatched_usage_types.items(),
                     key=lambda kv: kv[1], reverse=True)[:25]
        for name, amount in top:
            ws.cell(row=row, column=1, value=name)
            ws.cell(row=row, column=2, value=round(amount, 2)).number_format = MONEY
            row += 1

    if NOTES:
        row += 1
        ws.cell(row=row, column=1, value="Skips and warnings").font = TITLE_FONT
        row += 1
        write_header(ws, row, ["Area", "Detail"])
        row += 1
        for topic, detail in NOTES:
            ws.cell(row=row, column=1, value=topic)
            cell = ws.cell(row=row, column=2, value=detail)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
            ws.row_dimensions[row].height = 30
            row += 1

    finish_sheet(ws, [34, 30, 18, 18, 18, 18], account_id, header_row=1)


def build_workbook(path: str, spend: SpendData, vaults: List[Dict[str, Any]],
                   snapshots: List[Dict[str, Any]], rds: List[Dict[str, Any]],
                   ddb: List[Dict[str, Any]], redshift: List[Dict[str, Any]],
                   account_id: str, regions: Sequence[str], start: str, end: str,
                   months: int, generated: dt.datetime,
                   org: Optional[Dict[str, Any]] = None,
                   inventory_accounts: Optional[Sequence[str]] = None) -> None:
    wb = Workbook()
    sheet_summary(wb, spend, snapshots, account_id, start, end, months, org)
    if org:
        sheet_by_account(wb, org, account_id)
    sheet_monthly_detail(wb, spend, account_id)
    sheet_vaults(wb, vaults, account_id)
    sheet_snapshots(wb, snapshots, account_id)
    sheet_rds_dynamodb(wb, rds, ddb, redshift, account_id)
    sheet_notes(wb, spend, account_id, regions, start, end, generated, org,
                inventory_accounts)
    wb.save(path)


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Report what this AWS account actually pays for backup.")
    parser.add_argument("--profile", help="AWS profile name (default credential chain)")
    parser.add_argument("--months", type=int, default=12,
                        help="Months of Cost Explorer history (default 12, max 12)")
    parser.add_argument("--output", help="Output .xlsx path")
    parser.add_argument("--assume-role", metavar="ROLE_NAME",
                        help="Also gather the INVENTORY from every account in "
                             "the organization by assuming this role name in "
                             "each. Try OrganizationAccountAccessRole first; "
                             "it already exists in accounts created through "
                             "Organizations. Case-sensitive.")
    parser.add_argument("--single-account", action="store_true",
                        help="Skip the per-account breakdown even when run "
                             "from a management account")
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    months = max(1, min(args.months, 12))
    if months != args.months:
        note("Window", "--months %s clamped to %d." % (args.months, months))

    # Timezone-aware: utcnow() is deprecated on 3.12+, and boto3 returns
    # aware datetimes, so keeping everything aware makes the snapshot age
    # arithmetic straightforward.
    generated = dt.datetime.now(dt.timezone.utc)
    output = args.output or "aws-backup-costs-%s.xlsx" % generated.strftime("%Y-%m-%d")

    try:
        session = boto3.Session(profile_name=args.profile) if args.profile else boto3.Session()
        sts = session.client("sts", config=BOTO_CONFIG)
        account_id = sts.get_caller_identity()["Account"]
    except NoCredentialsError:
        print("No AWS credentials found. Configure a profile with `aws configure`,\n"
              "or set AWS_ACCESS_KEY_ID and AWS_SECRET_ACCESS_KEY.", file=sys.stderr)
        return 2
    except (ClientError, BotoCoreError) as err:
        print("Could not identify the AWS account: %s" % err, file=sys.stderr)
        return 2

    print("Account %s" % account_id)
    start, end = months_back(generated, months)

    print("Reading Cost Explorer ...")
    spend = fetch_spend(session, months, generated)

    org: Optional[Dict[str, Any]] = None
    if not args.single_account and spend.available:
        accounts = list_org_accounts(session)
        if accounts:
            print("Organization: %d account(s); querying Cost Explorer per "
                  "account (%d requests, about $%.2f) ..."
                  % (len(accounts), len(accounts), len(accounts) * 0.01))
            per_account = fetch_spend_by_account(session, months, generated,
                                                 accounts)
            rows = []
            unavailable = []
            for account in accounts:
                data = per_account[account["id"]]
                if not data.available:
                    unavailable.append(account)
                    continue
                rows.append({"id": account["id"], "name": account["name"],
                             "total": data.total,
                             "categories": data.category_totals()})
            rows.sort(key=lambda r: r["total"], reverse=True)
            org = {"rows": rows, "unavailable": unavailable}

    print("Listing enabled regions ...")
    regions = enabled_regions(session)
    print("  %d region(s)" % len(regions))

    inventory_accounts: List[Dict[str, str]] = []
    if args.assume_role and org:
        inventory_accounts = [{"id": e["id"], "name": e["name"]}
                              for e in org["rows"]]
        for entry in org["unavailable"]:
            inventory_accounts.append(entry)
    elif args.assume_role and not org:
        note("Organization inventory",
             "--assume-role was given but the organization could not be "
             "listed, so the inventory covers this account only.")

    vaults: List[Dict[str, Any]] = []
    snapshots: List[Dict[str, Any]] = []
    rds: List[Dict[str, Any]] = []
    ddb: List[Dict[str, Any]] = []
    redshift: List[Dict[str, Any]] = []
    inventory_scope_accounts = [account_id]

    if inventory_accounts:
        print("Gathering inventory across %d account(s) via role %s ..."
              % (len(inventory_accounts), args.assume_role))
        reached = []
        for entry in inventory_accounts:
            target = entry["id"]
            if target == account_id:
                # No point assuming a role into ourselves, and the role often
                # is not deployed to the management account at all — a
                # service-managed StackSet skips it.
                member = session
            else:
                member = assume_account_session(session, target, args.assume_role)
                if member is None:
                    print("  %s  skipped (see Notes)" % target)
                    continue
            print("  %s  %s" % (target, entry.get("name", "")))
            found = scan_account_inventory(member, regions, generated, target)
            vaults.extend(found["vaults"])
            snapshots.extend(found["snapshots"])
            rds.extend(found["rds"])
            ddb.extend(found["ddb"])
            redshift.extend(found["redshift"])
            reached.append(target)
        inventory_scope_accounts = reached
        note("Organization inventory",
             "Inventory gathered from %d of %d accounts using role %s."
             % (len(reached), len(inventory_accounts), args.assume_role))
    else:
        print("Scanning AWS Backup vaults ...")
        vaults = scan_vaults(session, regions)
        print("Scanning EBS snapshots ...")
        snapshots = scan_ebs_snapshots(session, regions, generated)
        print("Scanning RDS ...")
        rds = scan_rds(session, regions)
        print("Scanning DynamoDB ...")
        ddb = scan_dynamodb(session, regions)
        print("Scanning Redshift ...")
        redshift = scan_redshift(session, regions)
        for rows in (vaults, snapshots, rds, ddb, redshift):
            for row in rows:
                row["account"] = account_id

    build_workbook(output, spend, vaults, snapshots, rds, ddb, redshift,
                   account_id, regions, start, end, months, generated, org,
                   inventory_scope_accounts)

    print("")
    if spend.available:
        print("Total backup spend, %s to %s: %s %s"
              % (start, end, format(spend.total, ",.2f"), spend.currency))
        if org and org["rows"]:
            print("Across %d accounts:" % len(org["rows"]))
            for entry in org["rows"]:
                print("  %-14s %-28s %s"
                      % (entry["id"], entry["name"][:28],
                         format(entry["total"], ",.2f")))
        totals = dict((k, v) for k, v in spend.category_totals().items() if v)
        top3 = sorted(totals.items(), key=lambda kv: kv[1], reverse=True)[:3]
        if top3:
            print("Top categories:")
            for name, amount in top3:
                print("  %-24s %s" % (name, format(amount, ",.2f")))
    else:
        print("Total backup spend: UNAVAILABLE")
        print("  %s" % spend.error)

    orphans = [s for s in snapshots if is_orphan(s)]
    orphan_gib = sum(s.get("size_gib") or 0 for s in orphans)
    print("Orphaned EBS snapshots: %d (%s GiB)" % (len(orphans), format(orphan_gib, ",.0f")))
    print("Written: %s" % os.path.abspath(output))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
